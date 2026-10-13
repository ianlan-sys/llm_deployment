from pathlib import Path

import pytest

from app.services.review_parser import DocumentParseError, parse_document


def test_plain_text_parser_returns_sections(tmp_path):
    path = tmp_path / "proposal.txt"
    path.write_text(
        "一、项目范围\n缺少验收指标。\n二、实施计划\n计划为四周。",
        encoding="utf-8",
    )

    parsed = parse_document(path, "text/plain")

    assert [section.title for section in parsed.sections] == [
        "一、项目范围",
        "二、实施计划",
    ]
    assert parsed.sections[0].content == "一、项目范围\n缺少验收指标。"


def test_plain_text_parser_splits_oversized_sections(tmp_path):
    path = tmp_path / "proposal.txt"
    path.write_text("一、超长章节\n" + "内容" * 15, encoding="utf-8")

    parsed = parse_document(path, "text/plain", max_section_chars=20)

    assert len(parsed.sections) > 1
    assert all(len(section.content) <= 20 for section in parsed.sections)
    assert parsed.sections[0].title == "一、超长章节"


def test_plain_text_parser_reports_utf8_decode_failure(tmp_path):
    path = tmp_path / "proposal.txt"
    path.write_bytes(b"\xff\xfe\x00")

    with pytest.raises(DocumentParseError, match="UTF-8"):
        parse_document(path, "text/plain")


@pytest.mark.parametrize(
    ("filename", "media_type"),
    [
        ("proposal.txt", "application/pdf"),
        ("proposal.pdf", "text/plain"),
        ("proposal.doc", "application/msword"),
    ],
)
def test_parser_rejects_extension_and_media_type_mismatches(
    tmp_path, filename, media_type
):
    path = tmp_path / filename
    path.write_bytes(b"content")

    with pytest.raises(DocumentParseError, match="不支持|不匹配"):
        parse_document(path, media_type)


def test_pdf_parser_extracts_text_and_sections(tmp_path):
    path = tmp_path / "proposal.pdf"
    _write_text_pdf(path, ["1. Scope", "Acceptance metric missing."])

    parsed = parse_document(path, "application/pdf")

    assert parsed.sections[0].title == "1. Scope"
    assert "Acceptance metric missing." in parsed.sections[0].content


def test_docx_parser_reads_paragraphs(tmp_path):
    from docx import Document

    path = tmp_path / "proposal.docx"
    document = Document()
    document.add_paragraph("一、项目范围")
    document.add_paragraph("交付一套平台。")
    document.add_paragraph("二、实施计划")
    document.add_paragraph("计划为四周。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "里程碑"
    table.rows[0].cells[1].text = "第四周"
    document.save(path)

    parsed = parse_document(
        path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    assert [section.title for section in parsed.sections] == [
        "一、项目范围",
        "二、实施计划",
    ]
    assert "里程碑" in parsed.sections[1].content


def test_parser_rejects_documents_without_text(tmp_path):
    path = Path(tmp_path) / "empty.txt"
    path.write_text(" \n\t", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="没有可提取文本"):
        parse_document(path, "text/plain")


def _write_text_pdf(path: Path, lines: list[str]) -> None:
    escaped = [
        line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        for line in lines
    ]
    commands = ["BT /F1 12 Tf 72 720 Td"]
    for index, line in enumerate(escaped):
        if index:
            commands.append("0 -20 Td")
        commands.append(f"({line}) Tj")
    commands.append("ET")
    stream = " ".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode(
            "ascii"
        )
    )
    path.write_bytes(payload)
